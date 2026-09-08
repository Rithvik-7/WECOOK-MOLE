from pptx import Presentation

p = Presentation(r"C:\Users\brith\Downloads\SIH2026-IDEA-Presentation-Format_.pptx")
print("slides", len(p.slides))
print("w_in", p.slide_width.inches, "h_in", p.slide_height.inches)

for i, s in enumerate(p.slides, 1):
    print(f"\n=== SLIDE {i} layout={s.slide_layout.name} shapes={len(s.shapes)} ===")
    for j, sh in enumerate(s.shapes):
        txt = ""
        if getattr(sh, "has_text_frame", False):
            txt = sh.text_frame.text.replace("\n", " | ")[:400]
        print(
            f"  [{j}] {sh.name} type={sh.shape_type} "
            f"L={round(sh.left.inches,3)} T={round(sh.top.inches,3)} "
            f"W={round(sh.width.inches,3)} H={round(sh.height.inches,3)}"
        )
        print(f"      text={txt!r}")
        if sh.has_text_frame:
            for pi, para in enumerate(sh.text_frame.paragraphs):
                run_txt = " || ".join(r.text.replace("\n", "\\n") for r in para.runs if r.text)
                if run_txt:
                    sizes = []
                    for r in para.runs:
                        sizes.append(str(r.font.size.pt) if r.font.size else "?")
                    print(f"      p{pi} lvl={para.level} sizes={sizes[:8]} :: {run_txt[:250]}")
