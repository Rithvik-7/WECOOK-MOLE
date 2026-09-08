from pptx import Presentation

p = Presentation(r"C:\Users\brith\Desktop\mole\ppt\MOLE-SIH2026-Idea-Presentation.pptx")
print("slides", len(p.slides), "w", round(p.slide_width.inches, 2), "h", round(p.slide_height.inches, 2))
for i, s in enumerate(p.slides, 1):
    print(f"\n=== SLIDE {i} shapes={len(s.shapes)} ===")
    for j, sh in enumerate(s.shapes):
        txt = ""
        if getattr(sh, "has_text_frame", False):
            txt = sh.text_frame.text.replace("\n", " | ")[:180]
        print(
            f"  [{j}] {sh.name:22} L={sh.left.inches:5.2f} T={sh.top.inches:5.2f} "
            f"W={sh.width.inches:5.2f} H={sh.height.inches:5.2f}  {txt!r}"
        )
