"""Fill the official SIH 2026 6-slide idea template.

Rules we follow:
- Exactly 6 slides (delete the instructions slide).
- Keep SIH chrome, logo, footer '@SIH Idea submission- Template', slide numbers.
- Keep official section titles. Hide the giant instruction placeholders and replace
  them with diagrams / tables / infographics (SIH asks for points, not paragraphs).
- PS ID SIH26025 only. No fake accident %, collapse %, CO ppm, LoRa, autonomy, budget.
- Portal upload is PDF of these 6 slides.
"""
from __future__ import annotations

import shutil
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SRC = Path(r"C:\Users\brith\Downloads\SIH2026-IDEA-Presentation-Format_.pptx")
OUT = Path(r"C:\Users\brith\Desktop\mole\pitch\SIH26025_WE_COOK_IDEA.pptx")
OUT_DL = Path(r"C:\Users\brith\Downloads\SIH26025_WE_COOK_IDEA.pptx")
ASSETS = Path(r"C:\Users\brith\Desktop\mole\pitch\assets")

NAVY = RGBColor(0x0B, 0x2C, 0x4A)
ORANGE = RGBColor(0xF4, 0x7B, 0x20)
TEAL = RGBColor(0x0E, 0x7C, 0x7B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x4A, 0x55, 0x63)
PALE = RGBColor(0xF4, 0xF7, 0xFA)
LINE = RGBColor(0xD5, 0xDE, 0xE8)
GOLD = RGBColor(0xC4, 0x8A, 0x12)


def set_run(run, text, size=12, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def fill_lines(shape, lines, *, size=13, bold=False, color=INK, align=PP_ALIGN.LEFT, after=4):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(lines):
        if isinstance(item, tuple):
            text, st, bd, col = item
        else:
            text, st, bd, col = item, size, bold, color
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        set_run(p.add_run(), text, st, bd, col)


def hide_shape(shape):
    shape.left = Emu(0)
    shape.top = Emu(0)
    shape.width = Emu(1)
    shape.height = Emu(1)
    if shape.has_text_frame:
        shape.text_frame.clear()


def by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def no_line(shape):
    shape.line.fill.background()


def add_box(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.adjustments[0] = 0.08
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    no_line(sh)
    return sh


def add_tb(slide, l, t, w, h, lines, **kwargs):
    tb = slide.shapes.add_textbox(l, t, w, h)
    fill_lines(tb, lines, **kwargs)
    return tb


def set_cell_fill(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    for child in list(tc_pr):
        if child.tag == qn("a:solidFill"):
            tc_pr.remove(child)
    fill = etree.SubElement(tc_pr, qn("a:solidFill"))
    srgb = etree.SubElement(fill, qn("a:srgbClr"))
    srgb.set("val", hex_color)


def add_table(slide, l, t, w, h, rows, col_weights=None):
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), l, t, w, h)
    table = table_shape.table
    n_cols = len(rows[0])
    if col_weights is None:
        col_weights = [1] * n_cols
    total = sum(col_weights)
    for c, weight in enumerate(col_weights):
        table.columns[c].width = int(w * weight / total)
    row_h = int(h / len(rows))
    for r, row in enumerate(rows):
        table.rows[r].height = row_h
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            is_head = r == 0
            set_run(
                p.add_run(),
                val,
                size=11 if is_head else 11,
                bold=is_head or c == 0,
                color=WHITE if is_head else INK,
            )
            if is_head:
                set_cell_fill(cell, "0B2C4A")
            elif r % 2 == 0:
                set_cell_fill(cell, "F4F7FA")
            else:
                set_cell_fill(cell, "FFFFFF")
    return table_shape


def delete_slide(prs, index):
    sld_id = list(prs.slides._sldIdLst)[index]
    r_id = sld_id.get(qn("r:id"))
    prs.part.drop_rel(r_id)
    prs.slides._sldIdLst.remove(sld_id)


def set_team_oval(slide):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text or ""
        if "Your Team Name" in txt or txt.strip() == "WE COOK" or sh.name.startswith("Oval"):
            try:
                sh.fill.solid()
                sh.fill.fore_color.rgb = ORANGE
                no_line(sh)
            except Exception:
                pass
            fill_lines(sh, ["WE COOK"], size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER, after=0)
            sh.text_frame.anchor = MSO_ANCHOR.MIDDLE


def place_section_title(slide, text):
    sh = by_name(slide, "Title 1")
    if sh is None:
        return
    sh.left = Inches(1.88)
    sh.top = Inches(0.10)
    sh.width = Inches(8.55)
    sh.height = Inches(0.88)
    fill_lines(sh, [text], size=26, bold=True, color=NAVY, after=0)


def hide_instruction_box(slide):
    for sh in slide.shapes:
        if sh.name.startswith("TextBox") and sh.has_text_frame:
            t = sh.text_frame.text
            if any(
                key in t
                for key in (
                    "Proposed Solution",
                    "Technologies to be used",
                    "Analysis of the feasibility",
                    "Potential impact",
                    "Details / Links of the reference",
                )
            ):
                hide_shape(sh)


def fill_slide1(prs):
    s = prs.slides[0]
    sub = by_name(s, "Subtitle 3")
    if sub and sub.has_text_frame:
        fill_lines(
            sub,
            [
                ("TITLE PAGE", 20, True, NAVY),
                ("MOLE  ·  Mine Observation & Live-alert Engine", 16, False, ORANGE),
            ],
            after=2,
        )

    box = by_name(s, "TextBox 9")
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    fields = [
        ("Problem Statement ID", "SIH26025"),
        (
            "Problem Statement Title",
            "Development of an AI-enabled Low Cost Real Time Mine Subsidence Monitoring, Prediction and Early Warning System for Underground Coal Mines in India",
        ),
        ("Theme", "Disaster Management"),
        ("PS Category", "Hardware"),
        ("Organisation", "Ministry of Coal"),
        ("Team ID", "To be filled on the SIH portal after college shortlist"),
        ("Team Name", "WE COOK"),
        ("Idea", "MOLE — monitor the model, warn from evidence, inspect on command"),
    ]
    first = True
    for label, val in fields:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(7)
        set_run(p.add_run(), label + "   ", size=13, bold=True, color=ORANGE)
        set_run(p.add_run(), val, size=13, bold=False, color=NAVY)


def fill_slide2(prs):
    s = prs.slides[1]
    set_team_oval(s)
    hide_instruction_box(s)
    place_section_title(s, "MOLE — Proposed Solution")

    add_tb(
        s,
        Inches(0.35),
        Inches(1.08),
        Inches(12.6),
        Inches(0.32),
        ["Two fixed nodes watch the model. Laptop rules + sklearn explain. Rover inspects only after the officer decides."],
        size=13,
        bold=True,
        color=NAVY,
        after=0,
    )

    s.shapes.add_picture(str(ASSETS / "solution_flow.png"), Inches(0.35), Inches(1.42), Inches(12.6), Inches(2.22))

    add_tb(
        s,
        Inches(0.35),
        Inches(3.70),
        Inches(12.6),
        Inches(0.24),
        ["How this addresses SIH26025  ·  innovation = honest mapping, not a bigger claim"],
        size=12,
        bold=True,
        color=NAVY,
        after=0,
    )
    add_table(
        s,
        Inches(0.35),
        Inches(3.96),
        Inches(12.6),
        Inches(2.80),
        [
            ["PS word", "What MOLE actually does", "What we do not claim"],
            ["Monitoring", "Live Node A + Node B, source-tagged and timed", "Rover IMU is not Node B"],
            ["Early warning", "Rules: 3 samples → WATCH/ALERT; stale >5 s → UNKNOWN", "Green = healthy data, not a certified-safe mine"],
            ["AI-enabled", "Isolation Forest + LOF vs this rig’s NORMAL; score + top feature", "No LLM, no camera vision, no collapse %"],
            ["Prediction", "Holdout-selected Ridge: next 30 s of the same signals + MAE", "Not a mine-subsidence or roof-fall forecast"],
            ["Uniqueness", "ML cannot clear a rule latch. A vs B residual. Rover only after a warning.", "Not a gas-car maze robot; not LoRa; not autonomous"],
        ],
        col_weights=[1.15, 2.55, 2.4],
    )


def fill_slide3(prs):
    s = prs.slides[2]
    set_team_oval(s)
    hide_instruction_box(s)
    place_section_title(s, "TECHNICAL APPROACH")

    add_tb(
        s,
        Inches(0.35),
        Inches(1.02),
        Inches(12.6),
        Inches(0.26),
        ["Technologies + method  ·  two radios  ·  one local laptop  ·  rover is not step 6 of USB"],
        size=12,
        bold=True,
        color=NAVY,
        after=0,
    )

    tech = [
        (ASSETS / "icon_nodes.png", NAVY, "NODES A + B", "ESP32-WROOM-32 · MPU 0x68", "A: slider GPIO34   B: no slider"),
        (ASSETS / "icon_usb.png", NAVY, "RECEIVER", "Waveshare ESP32-S3-Zero", "ESP-NOW ch1 → USB JSON · no ML"),
        (ASSETS / "icon_ml.png", ORANGE, "LAPTOP ML", "Flask · SQLite · scikit-learn", "Isolation Forest + Ridge 30 s"),
        (ASSETS / "icon_inspect.png", TEAL, "ROVER", "L298N · HC-SR04 · IR GPIO19", "MQ-7 raw · AP Mine-Rover-AP"),
    ]
    x = Inches(0.35)
    card_w = Inches(3.05)
    for icon, accent, title, line1, line2 in tech:
        add_box(s, x, Inches(1.30), card_w, Inches(1.18), PALE)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.30), Inches(0.10), Inches(1.18))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        no_line(bar)
        s.shapes.add_picture(str(icon), x + Inches(0.18), Inches(1.40), Inches(0.38), Inches(0.38))
        add_tb(s, x + Inches(0.62), Inches(1.38), Inches(2.30), Inches(0.32), [(title, 13, True, accent)], after=0)
        add_tb(
            s,
            x + Inches(0.18),
            Inches(1.80),
            Inches(2.74),
            Inches(0.60),
            [(line1, 11, False, INK), (line2, 11, False, MUTED)],
            after=1,
        )
        x += card_w + Inches(0.15)

    s.shapes.add_picture(
        str(ASSETS / "architecture_lanes.png"),
        Inches(0.35),
        Inches(2.56),
        Inches(12.6),
        Inches(2.48),
    )

    engines = [
        (NAVY, "RULES  (latch)", "3 consecutive samples", "WATCH  3° or 2 mm", "ALERT  6° or 4 mm", "Stale >5 s → UNKNOWN"),
        (ORANGE, "ISOLATION FOREST + LOF", "Unusual vs this rig’s NORMAL", "Prior READY, or 120 live samples", "Score + top feature on screen", "Cannot clear a rule latch"),
        (TEAL, "30 s FORECAST", "Holdout-selected Ridge", "Next 30 s of same signals", "MAE shown on the chart", "Not a collapse prediction"),
    ]
    x = Inches(0.35)
    ew = Inches(4.08)
    for fill, title, a, b, c, dline in engines:
        add_box(s, x, Inches(5.12), ew, Inches(1.62), fill)
        add_tb(s, x + Inches(0.12), Inches(5.16), Inches(3.84), Inches(0.30), [(title, 13, True, WHITE)], after=0)
        add_tb(
            s,
            x + Inches(0.12),
            Inches(5.46),
            Inches(3.84),
            Inches(1.22),
            [
                (f"•  {a}", 12, False, WHITE),
                (f"•  {b}", 12, False, WHITE),
                (f"•  {c}", 12, False, WHITE),
                (f"•  {dline}", 12, False, WHITE),
            ],
            after=2,
        )
        x += ew + Inches(0.18)


def fill_slide4(prs):
    s = prs.slides[3]
    set_team_oval(s)
    hide_instruction_box(s)
    place_section_title(s, "FEASIBILITY AND VIABILITY")

    add_tb(
        s,
        Inches(0.35),
        Inches(1.02),
        Inches(12.6),
        Inches(0.26),
        ["We can run this on a table this week. It is not a real mine system."],
        size=14,
        bold=True,
        color=NAVY,
        after=0,
    )

    ready = [
        (ASSETS / "icon_ok.png", TEAL, "Parts are with us", "Node A, Node B, rover, USB board."),
        (ASSETS / "icon_dash.png", NAVY, "One laptop", "Local website. No login. No cloud."),
        (ASSETS / "icon_ml.png", ORANGE, "AI is ready", "Loaded already. No long wait."),
        (ASSETS / "icon_wifi.png", NAVY, "Short-range radio", "Works across the table. Not through rock."),
    ]
    x = Inches(0.35)
    cw = Inches(3.05)
    for icon, accent, title, body in ready:
        add_box(s, x, Inches(1.30), cw, Inches(1.08), PALE)
        s.shapes.add_picture(str(icon), x + Inches(0.10), Inches(1.40), Inches(0.42), Inches(0.42))
        add_tb(s, x + Inches(0.58), Inches(1.36), Inches(2.36), Inches(0.32), [(title, 14, True, accent)], after=0)
        add_tb(s, x + Inches(0.12), Inches(1.82), Inches(2.80), Inches(0.48), [(body, 12, False, INK)], after=0)
        x += cw + Inches(0.15)

    s.shapes.add_picture(str(ASSETS / "demo_loop.png"), Inches(0.35), Inches(2.46), Inches(12.6), Inches(1.52))

    s.shapes.add_picture(str(ASSETS / "threshold_graph.png"), Inches(0.35), Inches(4.06), Inches(6.15), Inches(2.62))

    fixes = [
        (ASSETS / "icon_stop.png", "Cable dies", "Show UNKNOWN. Never fake a green light."),
        (ASSETS / "icon_wifi.png", "Rover Wi-Fi dies", "Motors stop on their own in ~0.4 s."),
        (ASSETS / "icon_warn.png", "Easy to over-claim", "No collapse %. Tilt is not ground sink."),
        (ASSETS / "icon_inspect.png", "Toy electronics", "Table model only. Not mine-approved gear."),
    ]
    positions = [
        (Inches(6.65), Inches(4.06)),
        (Inches(9.85), Inches(4.06)),
        (Inches(6.65), Inches(5.42)),
        (Inches(9.85), Inches(5.42)),
    ]
    for (icon, title, body), (fx, fy) in zip(fixes, positions):
        add_box(s, fx, fy, Inches(3.10), Inches(1.26), PALE)
        s.shapes.add_picture(str(icon), fx + Inches(0.12), fy + Inches(0.14), Inches(0.42), Inches(0.42))
        add_tb(s, fx + Inches(0.62), fy + Inches(0.12), Inches(2.36), Inches(0.32), [(title, 13, True, NAVY)], after=0)
        add_tb(s, fx + Inches(0.12), fy + Inches(0.64), Inches(2.86), Inches(0.54), [(body, 12, False, INK)], after=0)


def fill_slide5(prs):
    s = prs.slides[4]
    set_team_oval(s)
    hide_instruction_box(s)
    place_section_title(s, "IMPACT AND BENEFITS")

    add_tb(
        s,
        Inches(0.35),
        Inches(1.08),
        Inches(12.6),
        Inches(0.40),
        ["Target on demo day: a mine safety officer on a tabletop model. Ministry of Coal owns the PS. Impact = notice + evidence + remote look — not disaster-prevention numbers we cannot prove."],
        size=13,
        bold=True,
        color=NAVY,
        after=0,
    )

    cards = [
        (NAVY, "Who", "Officer sees A vs B, rule reason, AI score, forecast MAE, then drives the rover because a warning appeared."),
        (TEAL, "Social", "Fewer unnecessary entries into the model tunnel. Decision stays human. Ack records awareness; it does not wipe history."),
        (ORANGE, "Economic / ops", "Standing ESP32-class watch vs waiting for a specialist survey. One laptop site. USB still works if venue Wi-Fi is bad."),
        (GOLD, "Environment / limit", "Hobby electronics on a model. Not intrinsically safe, not DGMS-approved, not a real underground vehicle."),
    ]
    x = Inches(0.35)
    for fill, title, body in cards:
        add_box(s, x, Inches(1.54), Inches(3.05), Inches(2.42), fill)
        add_tb(s, x + Inches(0.12), Inches(1.62), Inches(2.80), Inches(0.34), [title], size=14, bold=True, color=WHITE, after=0)
        add_tb(s, x + Inches(0.12), Inches(2.00), Inches(2.80), Inches(1.85), [body], size=12, color=WHITE, after=0)
        x += Inches(3.20)

    add_tb(
        s,
        Inches(0.35),
        Inches(4.08),
        Inches(12.6),
        Inches(0.24),
        ["Benefits a judge can verify on the table  ·  nothing we invent"],
        size=12,
        bold=True,
        color=NAVY,
        after=0,
    )
    add_table(
        s,
        Inches(0.35),
        Inches(4.34),
        Inches(12.6),
        Inches(2.42),
        [
            ["Benefit we can show", "How a judge verifies it", "Not a benefit we invent"],
            ["Earlier notice of disturbance", "Controlled tilt/slider → WATCH then latched ALERT", "Lives saved / disasters averted"],
            ["Evidence, not a vibe", "Charts + A-vs-B + IF feature + Ridge MAE", "87% collapse risk"],
            ["Inspect without replacing monitors", "Rover telemetry while A/B keep streaming", "Autonomous rescue robot"],
            ["Honest scale path", "Same workflow later with surveyed mounts", "This prototype already predicts subsidence"],
        ],
        col_weights=[1.5, 2.3, 2.0],
    )


def fill_slide6(prs):
    s = prs.slides[5]
    set_team_oval(s)
    hide_instruction_box(s)
    place_section_title(s, "RESEARCH AND REFERENCES")

    add_tb(
        s,
        Inches(0.35),
        Inches(1.08),
        Inches(12.6),
        Inches(0.32),
        ["References used to design the demo — not proof that MOLE is a certified mine instrument."],
        size=13,
        bold=True,
        color=NAVY,
        after=0,
    )
    add_table(
        s,
        Inches(0.35),
        Inches(1.42),
        Inches(12.6),
        Inches(4.55),
        [
            ["Item", "Why it is here", "Citation / link (no invented stats)"],
            ["PS SIH26025", "Official catalogue ID — never SIH2026025", "SIH 2026 portal · Ministry of Coal · Hardware · Disaster Management"],
            ["Isolation Forest", "Anomaly vs this rig’s normal", "Liu, Ting, Zhou — Isolation Forest, IEEE ICDM 2008"],
            ["30 s forecast", "Short-horizon prediction of measured signals", "scikit-learn Ridge / holdout MAE on a recent window"],
            ["MPU6050", "Gravity tilt + RMS vibration proxy", "InvenSense MPU-6000/6050 Product Specification"],
            ["ESP-NOW + ADC1", "Node radio ch1; slider on GPIO34", "Espressif ESP-NOW User Guide; ESP32 ADC notes"],
            ["ESP32-S3-Zero", "USB CDC receiver, no sensors", "Waveshare ESP32-S3-Zero wiki"],
            ["MQ-7 / IR", "Rover inspect only", "MQ-7 datasheet — raw ADC, not ppm. IR = digital LOW flag, not ranging"],
            ["Implementation", "Partial working software + firmware", "https://github.com/Rithvik-7/WECOOK-MOLE"],
            ["Domain context only", "Why mines care about movement", "Coal Mines Regulations, 2017; DGMS (S&T) Circ. 01/2017 cl. 6.1 — not implemented here"],
        ],
        col_weights=[1.35, 2.0, 3.1],
    )

    add_box(s, Inches(0.35), Inches(6.05), Inches(12.6), Inches(0.72), PALE)
    add_tb(
        s,
        Inches(0.48),
        Inches(6.10),
        Inches(12.35),
        Inches(0.62),
        [
            "Close: nodes detect → rules latch → Isolation Forest + Ridge explain → rover inspects.",
            "Upload PDF of these 6 slides only. Team ID: fill on portal. No extra slide.",
        ],
        size=12,
        bold=True,
        color=NAVY,
        after=2,
    )


def main():
    if not SRC.exists():
        raise SystemExit(f"Missing official template: {SRC}")
    missing = [
        "icon_nodes.png",
        "icon_warn.png",
        "icon_inspect.png",
        "icon_usb.png",
        "icon_ml.png",
        "solution_flow.png",
        "architecture_lanes.png",
        "icon_ok.png",
        "icon_stop.png",
        "icon_wifi.png",
        "icon_dash.png",
        "demo_loop.png",
        "threshold_graph.png",
    ]
    for name in missing:
        if not (ASSETS / name).exists():
            raise SystemExit(f"Run pitch/make_visuals.py first — missing {name}")

    prs = Presentation(str(SRC))
    assert len(prs.slides) == 7
    fill_slide1(prs)
    fill_slide2(prs)
    fill_slide3(prs)
    fill_slide4(prs)
    fill_slide5(prs)
    fill_slide6(prs)
    delete_slide(prs, 6)
    assert len(prs.slides) == 6
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    shutil.copyfile(OUT, OUT_DL)
    print("saved", OUT)
    print("copied", OUT_DL)
    print("slides", len(prs.slides))


if __name__ == "__main__":
    main()
