"""Fill the official SIH 2026 6-slide template. Do not add slides or rename official headings."""
from __future__ import annotations

import shutil
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(r"C:\Users\brith\Desktop\mole\ppt")
TEMPLATE = Path(r"C:\Users\brith\Downloads\SIH2026-IDEA-Presentation-Format_.pptx")
OUT = ROOT / "MOLE-SIH2026-Idea-Presentation.pptx"
OUT_DOWNLOADS = Path(r"C:\Users\brith\Downloads\MOLE-SIH2026-Idea-Presentation.pptx")

NAVY = RGBColor(0x1F, 0x49, 0x7D)
NAVY_DARK = RGBColor(0x14, 0x30, 0x54)
ORANGE = RGBColor(0xF7, 0x96, 0x46)
BLUE = RGBColor(0x4F, 0x81, 0xBD)
TEAL = RGBColor(0x4B, 0xAC, 0xC6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1C, 0x23, 0x2D)
MUTED = RGBColor(0x5A, 0x64, 0x70)
CARD = RGBColor(0xF4, 0xF7, 0xFB)


def delete_slide(prs: Presentation, index: int) -> None:
    sld_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(sld_id.rId)
    del prs.slides._sldIdLst[index]


def set_run_font(run, size=14, bold=False, color=INK, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def fill_text_frame(tf, lines, *, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, space_after=6):
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(lines):
        if isinstance(item, tuple) and len(item) == 4:
            text, st, bd, col = item
        else:
            text, st, bd, col = item, size, bold, color
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.level = 0
        run = p.add_run()
        run.text = text
        set_run_font(run, st, bd, col)


def shape_by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def set_shape_text(slide, name, lines, **kwargs):
    sh = shape_by_name(slide, name)
    if sh is None or not sh.has_text_frame:
        raise KeyError(name)
    fill_text_frame(sh.text_frame, lines, **kwargs)
    return sh


def set_simple(slide, name, text, size=14, bold=False, color=INK):
    return set_shape_text(slide, name, [text], size=size, bold=bold, color=color)


def remove_shape(slide, name):
    sh = shape_by_name(slide, name)
    if sh is None:
        return
    sh._element.getparent().remove(sh._element)


def hide_template_footer(slide):
    """Drop the official '@SIH Idea submission- Template' footer. Keep slide numbers."""
    remove_shape(slide, "Footer Placeholder 6")
    remove_shape(slide, "Footer Placeholder 5")
    # Bottom chrome bar only exists to carry that footer text.
    for name in ("Rectangle 8", "Rectangle 9"):
        remove_shape(slide, name)


def set_oval(slide, name, text="WE COOK"):
    sh = shape_by_name(slide, name)
    if sh is None or not sh.has_text_frame:
        return
    fill_text_frame(sh.text_frame, [text], size=12, bold=True, color=NAVY, align=PP_ALIGN.CENTER, space_after=0)
    sh.text_frame.anchor = MSO_ANCHOR.MIDDLE


def hide_or_compress_placeholder(slide, name, lines, top=Inches(1.18), height=Inches(0.32)):
    sh = shape_by_name(slide, name)
    if sh is None:
        return
    sh.top = top
    sh.left = Inches(0.40)
    sh.width = Inches(12.40)
    sh.height = height
    fill_text_frame(sh.text_frame, lines, size=12, bold=False, color=MUTED, space_after=0)


def no_line(shape):
    shape.line.fill.background()


def add_rect(slide, left, top, width, height, fill, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    no_line(sh)
    return sh


def add_textbox(slide, left, top, width, height, lines, **kwargs):
    sh = slide.shapes.add_textbox(left, top, width, height)
    fill_text_frame(sh.text_frame, lines, **kwargs)
    return sh


def add_bullets(slide, left, top, width, height, items, size=15, color=INK, space_after=8):
    lines = []
    for item in items:
        if isinstance(item, tuple):
            text, st, bd, col = item
            lines.append(("•  " + text, st, bd, col))
        else:
            lines.append(("•  " + item, size, False, color))
    return add_textbox(slide, left, top, width, height, lines, space_after=space_after)


def add_card(slide, left, top, width, height, fill, title, body_lines, title_size=14, body_size=13):
    add_rect(slide, left, top, width, height, fill, radius=True)
    lines = [(title, title_size, True, WHITE)]
    for b in body_lines:
        lines.append((b, body_size, False, WHITE))
    box = slide.shapes.add_textbox(left + Inches(0.12), top + Inches(0.10), width - Inches(0.24), height - Inches(0.18))
    fill_text_frame(box.text_frame, lines, space_after=4)
    return box


def add_light_card(slide, left, top, width, height, title, items, accent=ORANGE):
    add_rect(slide, left, top, width, height, CARD, radius=True)
    bar = add_rect(slide, left, top, Inches(0.10), height, accent)
    _ = bar
    add_textbox(
        slide,
        left + Inches(0.22),
        top + Inches(0.10),
        width - Inches(0.32),
        Inches(0.32),
        [(title, 15, True, NAVY)],
        space_after=0,
    )
    add_bullets(
        slide,
        left + Inches(0.18),
        top + Inches(0.42),
        width - Inches(0.30),
        height - Inches(0.50),
        items,
        size=13,
        space_after=6,
    )


def set_cell_fill(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    for child in list(tc_pr):
        if child.tag == qn("a:solidFill"):
            tc_pr.remove(child)
    fill = etree.SubElement(tc_pr, qn("a:solidFill"))
    srgb = etree.SubElement(fill, qn("a:srgbClr"))
    srgb.set("val", hex_color)


def build():
    shutil.copyfile(TEMPLATE, OUT)
    prs = Presentation(str(OUT))
    assert len(prs.slides) == 7
    delete_slide(prs, 6)
    assert len(prs.slides) == 6

    s1, s2, s3, s4, s5, s6 = prs.slides
    for s in (s1, s2, s3, s4, s5, s6):
        hide_template_footer(s)

    # --- SLIDE 1 ---
    set_shape_text(
        s1,
        "TextBox 9",
        [
            ("Problem Statement ID – SIH26025", 18, True, NAVY),
            (
                "Problem Statement Title – Development of an AI-enabled Low Cost Real Time Mine Subsidence Monitoring, Prediction and Early Warning System for Underground Coal Mines in India",
                13,
                False,
                INK,
            ),
            ("Theme – Disaster Management", 16, True, NAVY),
            ("PS Category – Hardware", 16, True, NAVY),
            ("Team ID – To be assigned on the SIH portal after college shortlist", 14, False, INK),
            ("Team Name – WE COOK", 16, True, NAVY),
        ],
        space_after=10,
    )
    sub = shape_by_name(s1, "Subtitle 3")
    fill_text_frame(
        sub.text_frame,
        [
            ("TITLE PAGE", 28, True, NAVY),
            ("MOLE  ·  Mine Observation & Live-alert Engine", 16, False, ORANGE),
        ],
        space_after=2,
    )

    # --- SLIDE 2 ---
    set_simple(s2, "Title 1", "MOLE", size=28, bold=True, color=NAVY)
    set_oval(s2, "Oval 9")
    hide_or_compress_placeholder(
        s2,
        "TextBox 8",
        ["Proposed solution  ·  how it addresses SIH26025  ·  what is unique"],
    )

    steps = [
        (NAVY, "1  Nodes A + B", "Watch continuously"),
        (BLUE, "2  Rules + AI", "Flag unusual change"),
        (ORANGE, "3  Dashboard", "Show the evidence"),
        (NAVY_DARK, "4  Operator", "Decide to inspect"),
        (TEAL, "5  Rover", "Look, then report"),
    ]
    box_w = Inches(2.20)
    gap = Inches(0.18)
    x0 = Inches(0.40)
    y0 = Inches(1.52)
    for i, (col, title, subtxt) in enumerate(steps):
        x = x0 + i * (box_w + gap)
        add_card(s2, x, y0, box_w, Inches(1.05), col, title, [subtxt], title_size=14, body_size=13)
        if i < 4:
            ar = s2.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                x + box_w + Inches(0.02),
                y0 + Inches(0.38),
                Inches(0.14),
                Inches(0.22),
            )
            ar.fill.solid()
            ar.fill.fore_color.rgb = ORANGE
            no_line(ar)

    add_textbox(
        s2,
        Inches(0.40),
        Inches(2.64),
        Inches(12.50),
        Inches(0.32),
        [
            (
                "Isolation Forest is early anomaly detection — a precursor stage toward prediction. It does not output a collapse probability.",
                14,
                True,
                NAVY,
            )
        ],
        space_after=0,
    )

    add_light_card(
        s2,
        Inches(0.40),
        Inches(3.02),
        Inches(6.10),
        Inches(2.35),
        "What MOLE does for SIH26025",
        [
            "Node A watches surface tilt vs a captured baseline.",
            "Node B watches vibration / disturbance at a second point.",
            "Watch / Alert latch with a written reason and history.",
            "Rover inspects only after the operator decides — it is not the monitor.",
            "One local dashboard. Rover vibration is never treated as Node B.",
        ],
        accent=NAVY,
    )
    add_light_card(
        s2,
        Inches(6.70),
        Inches(3.02),
        Inches(6.20),
        Inches(2.35),
        "Why this is not a copy of other student kits",
        [
            "Walk-in checks send a person in first. Nodes keep watching.",
            "InSAR / surveys are slow for a heading. This is a live tabletop loop.",
            "Rover-only kits have no standing surface watch.",
            "Unexplained AI scores hide the reason. Rules and AI are shown separately.",
        ],
        accent=ORANGE,
    )

    add_textbox(
        s2,
        Inches(0.40),
        Inches(5.44),
        Inches(12.50),
        Inches(1.30),
        [
            ("Scientific basis  (proxy, not a geophone)", 14, True, NAVY),
            (
                "•  MPU6050 tilt and vibration are low-cost proxies for mounting-surface micro-movement. They are not professional microseismic arrays.",
                13,
                False,
                INK,
            ),
            (
                "•  Published Indian longwall work: microseismic event-rate changes can precede roof / ground failure (Sivakumar et al., 2005, Rajendra / SECL).",
                13,
                False,
                INK,
            ),
            (
                "•  DGMS (S&T) Tech. Circular 01 of 2017, cl. 6.1 asks longwall mines for surface geophones. That is field-grade method — not this tabletop prototype.",
                13,
                False,
                INK,
            ),
            (
                "Stakes: CIL recorded 22 fatal accidents and 24 fatalities in 2024 up to November (MoC Annual Report 2024-25, Ch.14, Table 2; subject to DGMS reconciliation). Strata / roof control remains a named underground safety focus.",
                13,
                False,
                MUTED,
            ),
        ],
        space_after=4,
    )

    # --- SLIDE 3 ---
    set_oval(s3, "Oval 10")
    hide_or_compress_placeholder(
        s3,
        "TextBox 8",
        ["Technologies  ·  methodology  ·  working prototype"],
    )

    add_light_card(
        s3,
        Inches(0.40),
        Inches(1.52),
        Inches(6.10),
        Inches(2.65),
        "Hardware  (tabletop prototype)",
        [
            "Node A: ESP32-WROOM-32 + MPU6050 for tilt / movement.",
            "Node B: ESP32-WROOM-32 + MPU6050 for vibration.",
            "Rover: motors + driver, remote FWD / REV / LEFT / RIGHT / STOP.",
            "USB ESP32 gateway: radio/serial infrastructure, not a fourth node.",
            "Parts are already with the team. This is a demonstration rig, not mine-certified gear.",
        ],
        accent=NAVY,
    )
    add_light_card(
        s3,
        Inches(6.70),
        Inches(1.52),
        Inches(6.20),
        Inches(2.65),
        "Software  (one laptop)",
        [
            "Python dashboard. SQLite event history.",
            "Rules: Watch / Alert after 3 consecutive samples; 5 s stale → UNKNOWN.",
            "AI: Isolation Forest vs learned normal plus Ridge 30 s sensor forecast.",
            "AI cannot clear a rule alert. If untrained, show INACTIVE.",
            "Every reading is source-tagged and timed.",
        ],
        accent=ORANGE,
    )

    add_light_card(
        s3,
        Inches(0.40),
        Inches(4.28),
        Inches(6.10),
        Inches(2.55),
        "Communication",
        [
            "Prototype, surface nodes: ESP-NOW to the USB gateway.",
            "Rover commands use a separate control path. Not the node radio.",
            "Not claimed: underground Wi-Fi / ESP-NOW range through rock or coal.",
            "Field path (after an RF survey): LoRa, mesh, repeater, leaky-feeder, or mine-approved infrastructure — future work, not this demo.",
        ],
        accent=BLUE,
    )
    add_light_card(
        s3,
        Inches(6.70),
        Inches(4.28),
        Inches(6.20),
        Inches(2.55),
        "Method, every accepted reading",
        [
            "Is the packet fresh, valid, and from a known device ID?",
            "Compute change from the captured baseline (tilt / vibration).",
            "Demand persistence — one tap is not an alert.",
            "Run rules and Isolation Forest. Never let AI override a rule.",
            "Show the reason. If the officer decides, send the rover and log it.",
            "1 Hz firmware design. Latency and radio range: TBD — test pending.",
        ],
        accent=TEAL,
    )

    # --- SLIDE 4 ---
    set_oval(s4, "Oval 11")
    hide_or_compress_placeholder(
        s4,
        "TextBox 8",
        ["Feasibility  ·  challenges and risks  ·  how we handle them"],
    )

    add_card(
        s4,
        Inches(0.40),
        Inches(1.52),
        Inches(4.00),
        Inches(1.85),
        NAVY,
        "Build now",
        [
            "Two surface nodes, USB gateway, laptop dashboard, remote rover.",
            "Isolation Forest on live baseline samples.",
            "Local only — no cloud dependency for the demo.",
        ],
        title_size=16,
        body_size=13,
    )
    add_card(
        s4,
        Inches(4.60),
        Inches(1.52),
        Inches(4.00),
        Inches(1.85),
        ORANGE,
        "Honest limit",
        [
            "A tabletop model is not a mine.",
            "Tilt is orientation, not centimetres of settlement.",
            "Green means healthy data, not a certified-safe mine.",
        ],
        title_size=16,
        body_size=13,
    )
    add_card(
        s4,
        Inches(8.80),
        Inches(1.52),
        Inches(4.10),
        Inches(1.85),
        TEAL,
        "Field path",
        [
            "Surveyed mounts and site-calibrated thresholds.",
            "DGMS-approved / certified sensing and comms.",
            "Collapse models need geotech / InSAR labels we do not have.",
        ],
        title_size=16,
        body_size=13,
    )

    add_textbox(
        s4,
        Inches(0.40),
        Inches(3.48),
        Inches(12.50),
        Inches(0.30),
        [("Risks we already design for", 16, True, NAVY)],
        space_after=0,
    )
    add_bullets(
        s4,
        Inches(0.40),
        Inches(3.80),
        Inches(6.20),
        Inches(3.00),
        [
            "False green: a missing packet can look fine. After 5 s stale → UNKNOWN, never NORMAL.",
            "One noisy sample (tap, wire tug): require 3 consecutive samples before Watch / Alert.",
            "Rover IMU mistaken for Node B: same sensor family, different job. Device ID on every reading.",
            "ESP32 ADC2 fights wireless: crack slider only on GPIO34 / ADC1, or drop millimetre claims.",
        ],
        size=14,
        space_after=10,
    )
    add_bullets(
        s4,
        Inches(6.80),
        Inches(3.80),
        Inches(6.10),
        Inches(3.00),
        [
            "AI over-claim on a short tabletop set: Isolation Forest cannot clear a rule alert.",
            "Unlinked crack knob is not a gap: two-point ruler scale, or do not quote millimetres.",
            "Venue Wi-Fi / cloud demos die: laptop-local dashboard; backup video ready.",
            "COTS ESP32 / MPU6050 are not intrinsically safe. Tabletop only. Underground use needs DGMS-approved equipment for that hazardous area.",
        ],
        size=14,
        space_after=10,
    )

    # --- SLIDE 5 ---
    set_oval(s5, "Oval 11")
    hide_or_compress_placeholder(
        s5,
        "TextBox 8",
        ["Who it is for  ·  what we can honestly claim  ·  what comes next"],
    )

    add_light_card(
        s5,
        Inches(0.40),
        Inches(1.52),
        Inches(6.10),
        Inches(2.70),
        "Who this helps",
        [
            "Mine safety officer: evidence, written reason, and history — not a serial dump.",
            "Inspection crew: rover goes in after Watch / Alert; nodes keep watching.",
            "College / SIH judges: live loop — disturb → warn → inspect → log.",
            "Later mine operator: same dashboard language; certified sensors can replace hobby heads.",
        ],
        accent=NAVY,
    )
    add_light_card(
        s5,
        Inches(6.70),
        Inches(1.52),
        Inches(6.20),
        Inches(2.70),
        "Benefits we will actually stand behind",
        [
            "Social: fewer unnecessary entries into a disturbed heading on the demo analogue.",
            "Economic: a standing ESP32-class watch vs waiting for a specialist survey.",
            "Disaster: earlier notice of surface tilt and vibration — the PS proxy for subsidence, not proof of collapse.",
            "User experience: colour status, charts, written reason, rover pad. Acknowledging an alert does not erase the event.",
        ],
        accent=ORANGE,
    )

    add_textbox(
        s5,
        Inches(0.40),
        Inches(4.32),
        Inches(12.50),
        Inches(0.70),
        [
            ("Status discipline on the dashboard", 15, True, NAVY),
            (
                "NORMAL / WATCH / ALERT / UNKNOWN. Stale or missing telemetry → UNKNOWN, never false NORMAL. Green means healthy data. It does not mean a certified-safe mine. Rover is secondary inspection, idle until requested.",
                14,
                False,
                INK,
            ),
        ],
        space_after=4,
    )

    add_card(
        s5,
        Inches(0.40),
        Inches(5.12),
        Inches(4.00),
        Inches(1.55),
        NAVY,
        "NOW",
        ["Tabletop nodes + rover + local dashboard."],
        title_size=16,
        body_size=14,
    )
    add_card(
        s5,
        Inches(4.60),
        Inches(5.12),
        Inches(4.00),
        Inches(1.55),
        BLUE,
        "NEXT",
        ["More calibrated nodes, spatial view, extended validation."],
        title_size=16,
        body_size=14,
    )
    add_card(
        s5,
        Inches(8.80),
        Inches(5.12),
        Inches(4.10),
        Inches(1.55),
        ORANGE,
        "FIELD",
        [
            "Surveyed mounts, site thresholds, DGMS-approved sensing / comms, mine monitoring integration.",
        ],
        title_size=16,
        body_size=13,
    )

    # --- SLIDE 6 ---
    set_oval(s6, "Oval 8")
    hide_or_compress_placeholder(
        s6,
        "TextBox 8",
        ["References  ·  what we will show on 10 September"],
    )

    add_textbox(
        s6,
        Inches(0.40),
        Inches(1.50),
        Inches(12.50),
        Inches(3.55),
        [
            ("Sources actually used", 16, True, NAVY),
            ("1.  Coal Mines Regulations, 2017 · DGMS — duty to report dangerous occurrences.", 14, False, INK),
            ("2.  SIH 2026 · SIH26025 · Ministry of Coal · Hardware.", 14, False, INK),
            ("3.  MoC Annual Report 2024-25, Ch.14, Table 2 — CIL accident counts to Nov 2024 (not a cause-share %).", 14, False, INK),
            ("4.  Liu, Ting, Zhou — Isolation Forest, ICDM 2008. Library: scikit-learn IsolationForest.", 14, False, INK),
            ("5.  TDK / InvenSense MPU-6000/6050 datasheet — tilt / vibration sensor limits.", 14, False, INK),
            ("6.  Espressif ESP-NOW and ESP32 ADC docs — node radio; ADC1 vs wireless wiring.", 14, False, INK),
            ("7.  Sivakumar C., Srinivasan C., Gupta R.N. (2005) RaSiM 6 — Rajendra Colliery / SECL microseismic case.", 14, False, INK),
            ("8.  DGMS (S&T) Tech. Circular No. 01 of 2017, cl. 6.1 — longwall geophones; not implemented here.", 14, False, INK),
            ("9.  SBAS-InSAR literature — later fusion path only, not claimed as built.", 14, False, INK),
        ],
        space_after=5,
    )

    add_textbox(
        s6,
        Inches(0.40),
        Inches(5.10),
        Inches(12.50),
        Inches(1.60),
        [
            ("If a judge has seen bigger claims on other SIH decks", 15, True, NAVY),
            ("•  “AI predicts mine collapse %” → Isolation Forest on tabletop normal data + rule thresholds.", 14, False, INK),
            ("•  “LoRa mesh across a colliery” → ESP-NOW on the tabletop. LoRa is not demonstrated.", 14, False, INK),
            ("•  “Autonomous rover / AI vision” → remote drive. Camera only if shown live.", 14, False, INK),
            ("•  “Green = the mine is safe” → green = healthy data, no configured trigger.", 14, False, INK),
        ],
        space_after=4,
    )

    prs.save(str(OUT))
    shutil.copyfile(OUT, OUT_DOWNLOADS)
    print("saved", OUT)
    print("copied", OUT_DOWNLOADS)
    print("slides", len(prs.slides))


if __name__ == "__main__":
    build()
